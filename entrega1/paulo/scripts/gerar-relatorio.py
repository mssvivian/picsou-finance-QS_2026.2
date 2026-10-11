from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'entrega1/paulo/Relatorio-Paulo-Entrega1.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin, s.bottom_margin = Inches(.65), Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
normal = d.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
d.styles['Title'].font.size = Pt(26)
d.styles['Title'].paragraph_format.space_after = Pt(10)
d.styles['Heading 1'].font.size = Pt(15)
d.styles['Heading 1'].paragraph_format.space_before = Pt(13)
d.styles['Heading 1'].paragraph_format.space_after = Pt(7)
d.styles['Heading 2'].font.size = Pt(12)

header = s.header.paragraphs[0]
header.text = 'COS482   /   QUALIDADE DE SOFTWARE   /   ENTREGA 1'
header.runs[0].font.size = Pt(9)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Paulo Cesar Vaz  •  11 outubro 2026  |  ').font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text, style=None):
    return d.add_paragraph(text, style)

def h(text):
    d.add_heading(text, 1)

def path(text):
    z = p(text)
    z.paragraph_format.space_after = Pt(5)
    z.runs[0].font.name = 'Consolas'
    z.runs[0].font.size = Pt(9)

def table(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+edge)
        for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+key), val)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, width in enumerate(widths): t.columns[i].width = Inches(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i,text in enumerate(row): cells[i].text = text
    for rindex, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if rindex == 0: trpr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width = Inches(widths[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), '20364A' if rindex == 0 else ('F3F6F8' if rindex%2 else 'FFFFFF'));pr.append(shd)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge);m.set(qn('w:w'),'95');m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            for z in c.paragraphs:
                z.paragraph_format.space_after=Pt(2);z.paragraph_format.line_spacing=1.03
                for run in z.runs:
                    run.font.size=Pt(10)
                    if rindex==0: run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    p('').paragraph_format.space_after = Pt(0)

p('Relatório de testes de\namortização de empréstimos', 'Title')
p('Contribuição individual de Paulo Cesar Vaz', 'Subtitle')
p('Este relatório reúne a contribuição de Paulo à primeira entrega do Picsou Finance: seleção da classe, testes próprios, execução pela interface e problemas encontrados. Foram aprovados 31 testes unitários da classe e 15 passos de três casos pela interface; dois defeitos adicionais foram confirmados. A publicação no fork e a consolidação pelo grupo permanecem pendentes.')
h('1 Escopo e escolha da classe')
p('LoanAmortizationService calcula parcelas, capital amortizado, juros, seguro, custos e saldo em uma data de referência. A classe possui laço de amortização, decisões sobre datas e valores e fórmulas financeiras, atendendo à exigência de uma classe com algoritmo e complexidade razoável.')
p('Complexidade ciclomática: 14 no método compute(Debt, LocalDate), com 13 decisões; soma dos métodos explícitos 30. Convenção: 1 + decisões, incluindo condicionais, laços, ternários e operadores de curto-circuito. O maior método permanece com 14 mesmo sem contar esses operadores.')
h('2 Versão e ambiente testados')
p('Fork do grupo: github.com/mssvivian/picsou-finance-QS_2026.2. Branch local: feature/paulo-entrega1-amortizacao. Os commits locais ainda precisam ser publicados; as hashes abaixo distinguem a base do fork do build executado.')
path('Base do fork   6b020c76cf093b8995a68df12a32e8c5b21e9144\nBuild local    58c71b93544c27ff6f9b032dfe2c4df69fe01aa0')
p('Código de produção sem alterações entre essas versões. Aplicação construída localmente, backend real e demo desativado, em localhost:8082. Imagem picsou-paulo:entrega1-58c71b9-lf; adaptação operacional de CRLF para LF no entrypoint. Testes pela interface em 11/10/2026: Windows, Codex In-app Browser, versão não exposta, idioma francês.')
h('3 Responsabilidade individual')
p('Paulo é responsável pelos cenários unitários acrescentados, casos de empréstimos, evidências, relatos de defeitos e contribuição ao Plano de Teste IEEE 829. A execução técnica foi assistida pelo Codex sob solicitação de Paulo e está identificada nos registros. Este relatório também fornece os números e exemplos para os slides.')

d.add_page_break()
h('4 Testes unitários e evidências')
p('JUnit 5 e AssertJ, com instanciação direta do serviço e data de referência controlada. A classe não exige banco, contexto Spring ou Mockito. Nove testes herdados foram preservados; a contribuição acrescenta 22 execuções aprovadas, identificáveis pelo diff e pelos commits.')
p('Cenários próprios: parcela automática sem juros; ajuste do último pagamento com parcela explícita; limite de amortização ao saldo; seguro e tarifa; principal zero e valores opcionais nulos; datas ausentes, iguais ou invertidas; parcelas pagas e saldo antes, durante e após o contrato. Os esperados usam cálculos independentes e regras documentadas.')
table(['Execução','Resultado'], [
    ['Classe original','9 aprovados'],
    ['Classe ampliada','31 aprovados no total'],
    ['Reprodução dos defeitos','2 falhas de asserção confirmadas; ativação explícita'],
    ['Suíte completa do backend','1271 casos: 1216 aprovados, 20 erros e 35 ignorados'],
], [2.2,4.8])
p('Os 20 erros de infraestrutura também ocorreram no código original e não foram atribuídos à classe de empréstimos. Dos 35 ignorados, 33 já existiam; dois são as reproduções de defeitos, desativadas na execução padrão. A suíte ampla não foi aprovada globalmente.')
h('5 Casos pela interface e TestLink')
table(['Caso','Verificação principal','Resultado'], [
    ['dd-5  CT-PCV-01','1200 em 12 meses, juros zero: parcela 100, custo 1200 e saldo final zero','5 passos aprovados'],
    ['dd-6  CT-PCV-02','Parcela 110, capital 100, seguro total 120 e tarifa 50: custo 1370','5 passos aprovados'],
    ['dd-7  CT-PCV-03','Editar principal para 2400: parcela 200, persistência após recarga e sem duplicatas','5 passos aprovados'],
], [1.25,4.55,1.2])
p('Contratos futuros: zero parcelas pagas e saldo atual igual ao principal. Há 30 capturas originais distribuídas por teste, com ao menos uma por passo. A execução foi assistida pela interface real; não representa execução humana independente.')
p('TestLink: projeto Grupo DinDin; suíte Paulo - Amortização de empréstimos; plano Plano de Testes Manuais (46); build Paulo - Entrega1 - 58c71b9-lf (5). Execuções 8, 9 e 10 registradas como Passed, usando admin compartilhado e identificação do operador nas notas. Especificações em Draft para revisão. Parte dos anexos foi rejeitada pelo servidor; todas as imagens estão preservadas localmente.')

d.add_page_break()
h('6 Problemas encontrados')
table(['Registro local','Esperado','Observado'], [
    ['BUG-PCV-01\nArredondamento\nSeveridade proposta média','Principal 1000 em três parcelas automáticas: capitais somam 1000,00','Soma 999,99'],
    ['BUG-PCV-02\nSeguro na parcela automática\nSeveridade proposta alta','Principal 1200 em 12 meses, seguro mensal 10: parcela total 110,00','Parcela 100,00'],
], [2.15,3.75,1.1])
p('Os dois defeitos foram reproduzidos em testes unitários sobre o código original, com saída Maven e relatório Surefire. Os casos manuais aprovados usam divisões exatas e, no cenário com seguro, parcela explícita; portanto não exercitam essas falhas. Não há execução pela interface dos cenários específicos dos defeitos.')
p('Os relatos completos estão prontos, mas ainda precisam ser publicados no GitHub Issues do fork. Os identificadores BUG-PCV são locais. A entrega 1 exige registro em um bug tracker; arquivos locais e casos no TestLink não substituem essa publicação. Nenhuma regra de negócio foi corrigida nesta contribuição.')
h('7 Localização dos artefatos')
p('Caminhos relativos à raiz do repositório. O README principal aponta para os artefatos; entrega1/paulo/README.md reúne o índice individual.')
path('backend/src/test/java/com/picsou/service/\n  LoanAmortizationServiceTest.java\n  LoanAmortizationServiceDefectTest.java')
path('entrega1/paulo/\n  complexidade.md · matriz-de-testes.md · plano-de-teste.md\n  casos-manuais.md · resultados.md · ambiente-local.md\n  Tests_ScreenShots/README.md · Tests_ScreenShots/execucao.json\n  issues/README.md · issues/BUG-PCV-01.md · issues/BUG-PCV-02.md\n  evidencias/unitarios/ · evidencias/testlink/ · testlink/README.md')
h('8 Integração e preparação da apresentação')
p('Após receber acesso: publicar a branch preservando autoria; abrir PR; cadastrar as duas issues com evidências; integrar artefatos na main e conferir os links do README. Consolidar o Plano IEEE 829 no Google Docs com o grupo e verificar a submissão no Moodle até 11/10/2026 às 23h59, Brasília.')
p('Para os slides: explicar o algoritmo e a complexidade 14; apresentar os cenários próprios e os resultados; demonstrar a composição 1370 do contrato com seguro e tarifa; explicar os dois defeitos e as limitações. Integração, E2E, cobertura de arestas e mutação pertencem à segunda entrega e não foram avaliados aqui.')
d.core_properties.title = 'Relatório de testes de amortização de empréstimos'
d.core_properties.author = 'Paulo Cesar Vaz'
d.core_properties.subject = 'COS482 Entrega 1'
d.save(OUT)
print(OUT)
